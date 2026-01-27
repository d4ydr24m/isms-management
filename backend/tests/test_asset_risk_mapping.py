"""
자산-위협-취약점 3-way 매핑 테스트
Phase 2: FR-603, Section 5.1

TDD 방식:
1. 테스트 먼저 작성 (RED)
2. 최소 구현 (GREEN)
3. 리팩토링 (IMPROVE)

테스트 범위:
- 5.1.1: 자산 유형별 위협/취약점 자동 추천
- 5.1.2: 매핑 유효성 검증 로직
- 5.1.3: 대량 매핑 기능 (엑셀 임포트)
"""
import pytest
from datetime import date
from io import BytesIO
from typing import List

from sqlalchemy.orm import Session

from app.models.asset import Asset, AssetType
from app.models.risk import (
    Threat, Vulnerability, ThreatCategory, VulnerabilityCategory,
    AssetTypeThreat, RiskScenario, RiskAssessment,
)
from app.services.asset_risk_mapping_service import AssetRiskMappingService


# =============================================================================
# 픽스처
# =============================================================================

@pytest.fixture
def mapping_service(db: Session) -> AssetRiskMappingService:
    """AssetRiskMappingService 인스턴스"""
    return AssetRiskMappingService(db)


@pytest.fixture
def asset_type_server(db: Session) -> AssetType:
    """서버 자산 유형 생성"""
    asset_type = AssetType(
        code="SRV",
        name="서버",
        description="서버 시스템",
        is_custom=False,
        is_active=True,
        sort_order=1,
    )
    db.add(asset_type)
    db.commit()
    db.refresh(asset_type)
    return asset_type


@pytest.fixture
def asset_type_network(db: Session) -> AssetType:
    """네트워크 장비 자산 유형 생성"""
    asset_type = AssetType(
        code="NET",
        name="네트워크장비",
        description="네트워크 장비",
        is_custom=False,
        is_active=True,
        sort_order=2,
    )
    db.add(asset_type)
    db.commit()
    db.refresh(asset_type)
    return asset_type


@pytest.fixture
def threat_category(db: Session) -> ThreatCategory:
    """위협 분류 생성"""
    category = ThreatCategory(
        code="TC-SYS",
        name="시스템 위협",
        description="시스템 관련 위협",
        sort_order=1,
        is_active=True,
    )
    db.add(category)
    db.commit()
    db.refresh(category)
    return category


@pytest.fixture
def vulnerability_category(db: Session) -> VulnerabilityCategory:
    """취약점 분류 생성"""
    category = VulnerabilityCategory(
        code="VC-SYS",
        name="시스템 취약점",
        description="시스템 관련 취약점",
        sort_order=1,
        is_active=True,
    )
    db.add(category)
    db.commit()
    db.refresh(category)
    return category


@pytest.fixture
def server_threats(
    db: Session, threat_category: ThreatCategory, asset_type_server: AssetType
) -> List[Threat]:
    """서버용 위협 목록 생성"""
    threats = []
    threat_data = [
        ("T-SRV-001", "비인가 접근", "서버에 대한 비인가 접근 시도", 3),
        ("T-SRV-002", "악성코드 감염", "서버 악성코드 감염", 3),
        ("T-SRV-003", "서비스 거부 공격", "DoS/DDoS 공격", 2),
    ]

    for code, name, desc, level in threat_data:
        threat = Threat(
            code=code,
            name=name,
            description=desc,
            category_id=threat_category.id,
            threat_level=level,
            is_custom=False,
            is_active=True,
        )
        db.add(threat)
        db.flush()

        # 자산 유형-위협 매핑
        mapping = AssetTypeThreat(
            asset_type_id=asset_type_server.id,
            threat_id=threat.id,
            relevance_score=0.9,
            is_active=True,
        )
        db.add(mapping)
        threats.append(threat)

    db.commit()
    for t in threats:
        db.refresh(t)
    return threats


@pytest.fixture
def network_threats(
    db: Session, threat_category: ThreatCategory, asset_type_network: AssetType
) -> List[Threat]:
    """네트워크 장비용 위협 목록 생성"""
    threats = []
    threat_data = [
        ("T-NET-001", "패킷 스니핑", "네트워크 패킷 도청", 3),
        ("T-NET-002", "라우팅 공격", "라우팅 테이블 변조", 2),
    ]

    for code, name, desc, level in threat_data:
        threat = Threat(
            code=code,
            name=name,
            description=desc,
            category_id=threat_category.id,
            threat_level=level,
            is_custom=False,
            is_active=True,
        )
        db.add(threat)
        db.flush()

        # 자산 유형-위협 매핑
        mapping = AssetTypeThreat(
            asset_type_id=asset_type_network.id,
            threat_id=threat.id,
            relevance_score=0.85,
            is_active=True,
        )
        db.add(mapping)
        threats.append(threat)

    db.commit()
    for t in threats:
        db.refresh(t)
    return threats


@pytest.fixture
def server_vulnerabilities(
    db: Session, vulnerability_category: VulnerabilityCategory
) -> List[Vulnerability]:
    """서버용 취약점 목록 생성"""
    vulns = []
    vuln_data = [
        ("V-SRV-001", "패치 미적용", "보안 패치 미적용 취약점", 3),
        ("V-SRV-002", "약한 암호 설정", "취약한 암호 정책", 2),
        ("V-SRV-003", "불필요한 서비스", "불필요한 서비스 구동", 2),
    ]

    for code, name, desc, level in vuln_data:
        vuln = Vulnerability(
            code=code,
            name=name,
            description=desc,
            category_id=vulnerability_category.id,
            vulnerability_level=level,
            is_custom=False,
            is_active=True,
        )
        db.add(vuln)
        vulns.append(vuln)

    db.commit()
    for v in vulns:
        db.refresh(v)
    return vulns


@pytest.fixture
def sample_asset(
    db: Session, asset_type_server: AssetType, test_department, test_user
) -> Asset:
    """테스트용 자산 생성"""
    asset = Asset(
        asset_code="AST-SRV-202501-001",
        name="웹서버-01",
        description="테스트용 웹서버",
        asset_type_id=asset_type_server.id,
        department_id=test_department.id,
        owner_id=test_user.id,
        location="서버실 A",
        ip_address="192.168.1.10",
        status="운영",
        is_active=True,
    )
    db.add(asset)
    db.commit()
    db.refresh(asset)
    return asset


@pytest.fixture
def sample_scenario(db: Session, test_user) -> RiskScenario:
    """테스트용 위험 평가 시나리오"""
    scenario = RiskScenario(
        name="2025년 1차 위험평가",
        description="테스트 시나리오",
        start_date=date.today(),
        status="in_progress",
        created_by=test_user.id,
    )
    db.add(scenario)
    db.commit()
    db.refresh(scenario)
    return scenario


# =============================================================================
# 5.1.1: 자산 유형별 위협/취약점 자동 추천 테스트
# =============================================================================

class TestGetRecommendedThreats:
    """5.1.1 - 자산 유형별 위협 자동 추천 테스트"""

    def test_returns_threats_for_server_type(
        self, mapping_service, asset_type_server, server_threats
    ):
        """서버 유형에 대해 관련 위협을 반환한다"""
        # When
        result = mapping_service.get_recommended_threats(asset_type_server.id)

        # Then
        assert len(result) == 3
        threat_codes = [t.code for t in result]
        assert "T-SRV-001" in threat_codes
        assert "T-SRV-002" in threat_codes
        assert "T-SRV-003" in threat_codes

    def test_returns_threats_for_network_type(
        self, mapping_service, asset_type_network, network_threats
    ):
        """네트워크 유형에 대해 관련 위협을 반환한다"""
        # When
        result = mapping_service.get_recommended_threats(asset_type_network.id)

        # Then
        assert len(result) == 2
        threat_codes = [t.code for t in result]
        assert "T-NET-001" in threat_codes
        assert "T-NET-002" in threat_codes

    def test_returns_empty_for_type_without_mappings(
        self, db: Session, mapping_service
    ):
        """매핑되지 않은 자산 유형에 대해 빈 목록 반환"""
        # Given
        new_type = AssetType(
            code="OTHER",
            name="기타",
            is_custom=True,
            is_active=True,
        )
        db.add(new_type)
        db.commit()
        db.refresh(new_type)

        # When
        result = mapping_service.get_recommended_threats(new_type.id)

        # Then
        assert len(result) == 0

    def test_returns_empty_for_invalid_type_id(self, mapping_service):
        """존재하지 않는 자산 유형 ID에 대해 빈 목록 반환"""
        # When
        result = mapping_service.get_recommended_threats(99999)

        # Then
        assert len(result) == 0

    def test_ordered_by_relevance_score(
        self, db: Session, mapping_service, asset_type_server, threat_category
    ):
        """관련성 점수 순으로 정렬된 위협 반환"""
        # Given - 다른 관련성 점수로 위협 추가
        threat_high = Threat(
            code="T-HIGH",
            name="높은 관련성 위협",
            category_id=threat_category.id,
            threat_level=3,
            is_custom=False,
            is_active=True,
        )
        db.add(threat_high)
        db.flush()

        mapping_high = AssetTypeThreat(
            asset_type_id=asset_type_server.id,
            threat_id=threat_high.id,
            relevance_score=1.0,  # 가장 높음
            is_active=True,
        )
        db.add(mapping_high)

        threat_low = Threat(
            code="T-LOW",
            name="낮은 관련성 위협",
            category_id=threat_category.id,
            threat_level=1,
            is_custom=False,
            is_active=True,
        )
        db.add(threat_low)
        db.flush()

        mapping_low = AssetTypeThreat(
            asset_type_id=asset_type_server.id,
            threat_id=threat_low.id,
            relevance_score=0.3,  # 가장 낮음
            is_active=True,
        )
        db.add(mapping_low)
        db.commit()

        # When
        result = mapping_service.get_recommended_threats(asset_type_server.id)

        # Then
        assert result[0].code == "T-HIGH"  # 가장 높은 관련성이 먼저
        assert result[-1].code == "T-LOW"  # 가장 낮은 관련성이 마지막


class TestGetRecommendedVulnerabilities:
    """5.1.1 - 자산 유형별 취약점 자동 추천 테스트"""

    def test_returns_vulnerabilities_for_asset_type(
        self, mapping_service, asset_type_server, server_vulnerabilities
    ):
        """자산 유형에 대해 관련 취약점을 반환한다"""
        # When
        result = mapping_service.get_recommended_vulnerabilities(asset_type_server.id)

        # Then
        # 현재 취약점은 자산 유형에 직접 매핑되지 않으므로 전체 활성 취약점 반환
        assert len(result) >= 3

    def test_returns_only_active_vulnerabilities(
        self, db: Session, mapping_service, vulnerability_category, asset_type_server
    ):
        """비활성 취약점은 제외된다"""
        # Given
        inactive_vuln = Vulnerability(
            code="V-INACTIVE",
            name="비활성 취약점",
            category_id=vulnerability_category.id,
            vulnerability_level=1,
            is_custom=False,
            is_active=False,  # 비활성
        )
        db.add(inactive_vuln)
        db.commit()

        # When
        result = mapping_service.get_recommended_vulnerabilities(asset_type_server.id)

        # Then
        vuln_codes = [v.code for v in result]
        assert "V-INACTIVE" not in vuln_codes


# =============================================================================
# 5.1.2: 매핑 유효성 검증 로직 테스트
# =============================================================================

class TestValidateMapping:
    """5.1.2 - 매핑 유효성 검증 테스트"""

    def test_valid_mapping_returns_success(
        self,
        mapping_service,
        sample_asset,
        server_threats,
        server_vulnerabilities,
    ):
        """유효한 매핑은 성공 결과를 반환한다"""
        # When
        result = mapping_service.validate_mapping(
            asset_id=sample_asset.id,
            threat_id=server_threats[0].id,
            vulnerability_id=server_vulnerabilities[0].id,
        )

        # Then
        assert result["is_valid"] is True
        assert result["errors"] == []

    def test_invalid_asset_returns_error(
        self,
        mapping_service,
        server_threats,
        server_vulnerabilities,
    ):
        """존재하지 않는 자산 ID는 에러를 반환한다"""
        # When
        result = mapping_service.validate_mapping(
            asset_id=99999,
            threat_id=server_threats[0].id,
            vulnerability_id=server_vulnerabilities[0].id,
        )

        # Then
        assert result["is_valid"] is False
        assert "자산" in result["errors"][0]

    def test_invalid_threat_returns_error(
        self,
        mapping_service,
        sample_asset,
        server_vulnerabilities,
    ):
        """존재하지 않는 위협 ID는 에러를 반환한다"""
        # When
        result = mapping_service.validate_mapping(
            asset_id=sample_asset.id,
            threat_id=99999,
            vulnerability_id=server_vulnerabilities[0].id,
        )

        # Then
        assert result["is_valid"] is False
        assert "위협" in result["errors"][0]

    def test_invalid_vulnerability_returns_error(
        self,
        mapping_service,
        sample_asset,
        server_threats,
    ):
        """존재하지 않는 취약점 ID는 에러를 반환한다"""
        # When
        result = mapping_service.validate_mapping(
            asset_id=sample_asset.id,
            threat_id=server_threats[0].id,
            vulnerability_id=99999,
        )

        # Then
        assert result["is_valid"] is False
        assert "취약점" in result["errors"][0]

    def test_inactive_asset_returns_error(
        self,
        db: Session,
        mapping_service,
        sample_asset,
        server_threats,
        server_vulnerabilities,
    ):
        """비활성 자산은 에러를 반환한다"""
        # Given
        sample_asset.is_active = False
        db.commit()

        # When
        result = mapping_service.validate_mapping(
            asset_id=sample_asset.id,
            threat_id=server_threats[0].id,
            vulnerability_id=server_vulnerabilities[0].id,
        )

        # Then
        assert result["is_valid"] is False
        assert "비활성" in result["errors"][0]

    def test_inactive_threat_returns_error(
        self,
        db: Session,
        mapping_service,
        sample_asset,
        server_threats,
        server_vulnerabilities,
    ):
        """비활성 위협은 에러를 반환한다"""
        # Given
        server_threats[0].is_active = False
        db.commit()

        # When
        result = mapping_service.validate_mapping(
            asset_id=sample_asset.id,
            threat_id=server_threats[0].id,
            vulnerability_id=server_vulnerabilities[0].id,
        )

        # Then
        assert result["is_valid"] is False
        assert "비활성" in result["errors"][0]

    def test_returns_warning_for_unrelated_threat(
        self,
        mapping_service,
        sample_asset,
        network_threats,  # 서버 자산에 네트워크 위협
        server_vulnerabilities,
    ):
        """자산 유형과 관련 없는 위협에 대해 경고를 반환한다"""
        # When
        result = mapping_service.validate_mapping(
            asset_id=sample_asset.id,
            threat_id=network_threats[0].id,  # 네트워크 위협을 서버에 매핑
            vulnerability_id=server_vulnerabilities[0].id,
        )

        # Then
        # 유효하지만 경고 포함
        assert result["is_valid"] is True
        assert len(result["warnings"]) > 0
        assert "자산 유형" in result["warnings"][0]


class TestValidateMappingBatch:
    """5.1.2 - 일괄 매핑 유효성 검증 테스트"""

    def test_validates_multiple_mappings(
        self,
        mapping_service,
        sample_asset,
        server_threats,
        server_vulnerabilities,
    ):
        """여러 매핑을 일괄 검증한다"""
        # Given
        mappings = [
            {
                "asset_id": sample_asset.id,
                "threat_id": server_threats[0].id,
                "vulnerability_id": server_vulnerabilities[0].id,
            },
            {
                "asset_id": sample_asset.id,
                "threat_id": server_threats[1].id,
                "vulnerability_id": server_vulnerabilities[1].id,
            },
        ]

        # When
        result = mapping_service.validate_mapping_batch(mappings)

        # Then
        assert result["total"] == 2
        assert result["valid_count"] == 2
        assert result["invalid_count"] == 0

    def test_batch_with_invalid_mappings(
        self,
        mapping_service,
        sample_asset,
        server_threats,
        server_vulnerabilities,
    ):
        """유효하지 않은 매핑 포함 시 개수를 반환한다"""
        # Given
        mappings = [
            {
                "asset_id": sample_asset.id,
                "threat_id": server_threats[0].id,
                "vulnerability_id": server_vulnerabilities[0].id,
            },
            {
                "asset_id": 99999,  # 잘못된 자산
                "threat_id": server_threats[0].id,
                "vulnerability_id": server_vulnerabilities[0].id,
            },
        ]

        # When
        result = mapping_service.validate_mapping_batch(mappings)

        # Then
        assert result["total"] == 2
        assert result["valid_count"] == 1
        assert result["invalid_count"] == 1


# =============================================================================
# 5.1.3: 대량 매핑 기능 (엑셀 임포트) 테스트
# =============================================================================

class TestBulkCreateMappingsFromExcel:
    """5.1.3 - 엑셀 임포트를 통한 대량 매핑 테스트"""

    def test_creates_mappings_from_valid_excel(
        self,
        mapping_service,
        sample_asset,
        sample_scenario,
        server_threats,
        server_vulnerabilities,
        test_user,
    ):
        """유효한 엑셀 파일로 매핑을 생성한다"""
        # Given
        excel_content = self._create_valid_excel_file(
            sample_asset.asset_code,
            server_threats[0].code,
            server_vulnerabilities[0].code,
        )

        # When
        result = mapping_service.bulk_create_mappings_from_excel(
            file_content=excel_content,
            scenario_id=sample_scenario.id,
            user_id=test_user.id,
        )

        # Then
        assert result["success"] >= 1
        assert result["failed"] == 0

    def test_returns_errors_for_invalid_rows(
        self,
        mapping_service,
        sample_asset,
        sample_scenario,
        server_threats,
        server_vulnerabilities,
        test_user,
    ):
        """잘못된 행이 포함된 엑셀 파일 처리"""
        # Given
        excel_content = self._create_excel_with_invalid_rows()

        # When
        result = mapping_service.bulk_create_mappings_from_excel(
            file_content=excel_content,
            scenario_id=sample_scenario.id,
            user_id=test_user.id,
        )

        # Then
        assert result["failed"] > 0
        assert len(result["errors"]) > 0

    def test_returns_summary_with_row_details(
        self,
        mapping_service,
        sample_asset,
        sample_scenario,
        server_threats,
        server_vulnerabilities,
        test_user,
    ):
        """결과에 행별 상세 정보가 포함된다"""
        # Given
        excel_content = self._create_valid_excel_file(
            sample_asset.asset_code,
            server_threats[0].code,
            server_vulnerabilities[0].code,
        )

        # When
        result = mapping_service.bulk_create_mappings_from_excel(
            file_content=excel_content,
            scenario_id=sample_scenario.id,
            user_id=test_user.id,
        )

        # Then
        assert "total" in result
        assert "success" in result
        assert "failed" in result
        assert "errors" in result

    def test_rejects_empty_file(
        self,
        mapping_service,
        sample_scenario,
        test_user,
    ):
        """빈 파일은 에러를 반환한다"""
        # Given
        excel_content = self._create_empty_excel_file()

        # When
        result = mapping_service.bulk_create_mappings_from_excel(
            file_content=excel_content,
            scenario_id=sample_scenario.id,
            user_id=test_user.id,
        )

        # Then
        assert result["total"] == 0
        assert result["success"] == 0

    def test_skips_duplicate_mappings(
        self,
        db: Session,
        mapping_service,
        sample_asset,
        sample_scenario,
        server_threats,
        server_vulnerabilities,
        test_user,
    ):
        """중복 매핑은 건너뛴다"""
        # Given - 기존 매핑 생성
        existing = RiskAssessment(
            scenario_id=sample_scenario.id,
            asset_id=sample_asset.id,
            threat_id=server_threats[0].id,
            vulnerability_id=server_vulnerabilities[0].id,
            asset_value=2,
            threat_level=2,
            vulnerability_level=2,
        )
        db.add(existing)
        db.commit()

        excel_content = self._create_valid_excel_file(
            sample_asset.asset_code,
            server_threats[0].code,
            server_vulnerabilities[0].code,
        )

        # When
        result = mapping_service.bulk_create_mappings_from_excel(
            file_content=excel_content,
            scenario_id=sample_scenario.id,
            user_id=test_user.id,
        )

        # Then
        assert result["skipped"] >= 1

    # 헬퍼 메서드
    def _create_valid_excel_file(
        self, asset_code: str, threat_code: str, vuln_code: str
    ) -> bytes:
        """유효한 테스트용 엑셀 파일 생성"""
        try:
            from openpyxl import Workbook
        except ImportError:
            pytest.skip("openpyxl이 설치되지 않았습니다")

        wb = Workbook()
        ws = wb.active
        ws.title = "매핑"

        # 헤더
        headers = ["자산코드", "위협코드", "취약점코드", "자산가치", "위협등급", "취약점등급", "비고"]
        for col, header in enumerate(headers, 1):
            ws.cell(row=1, column=col, value=header)

        # 데이터
        ws.cell(row=2, column=1, value=asset_code)
        ws.cell(row=2, column=2, value=threat_code)
        ws.cell(row=2, column=3, value=vuln_code)
        ws.cell(row=2, column=4, value=2)  # 자산가치
        ws.cell(row=2, column=5, value=2)  # 위협등급
        ws.cell(row=2, column=6, value=2)  # 취약점등급
        ws.cell(row=2, column=7, value="테스트 매핑")

        output = BytesIO()
        wb.save(output)
        output.seek(0)
        return output.read()

    def _create_excel_with_invalid_rows(self) -> bytes:
        """잘못된 행이 포함된 엑셀 파일 생성"""
        try:
            from openpyxl import Workbook
        except ImportError:
            pytest.skip("openpyxl이 설치되지 않았습니다")

        wb = Workbook()
        ws = wb.active
        ws.title = "매핑"

        # 헤더
        headers = ["자산코드", "위협코드", "취약점코드", "자산가치", "위협등급", "취약점등급"]
        for col, header in enumerate(headers, 1):
            ws.cell(row=1, column=col, value=header)

        # 잘못된 데이터 (존재하지 않는 코드)
        ws.cell(row=2, column=1, value="INVALID-ASSET")
        ws.cell(row=2, column=2, value="INVALID-THREAT")
        ws.cell(row=2, column=3, value="INVALID-VULN")
        ws.cell(row=2, column=4, value=2)
        ws.cell(row=2, column=5, value=2)
        ws.cell(row=2, column=6, value=2)

        output = BytesIO()
        wb.save(output)
        output.seek(0)
        return output.read()

    def _create_empty_excel_file(self) -> bytes:
        """빈 엑셀 파일 생성 (헤더만)"""
        try:
            from openpyxl import Workbook
        except ImportError:
            pytest.skip("openpyxl이 설치되지 않았습니다")

        wb = Workbook()
        ws = wb.active
        ws.title = "매핑"

        # 헤더만
        headers = ["자산코드", "위협코드", "취약점코드", "자산가치", "위협등급", "취약점등급"]
        for col, header in enumerate(headers, 1):
            ws.cell(row=1, column=col, value=header)

        output = BytesIO()
        wb.save(output)
        output.seek(0)
        return output.read()


class TestGetMappingTemplate:
    """매핑 템플릿 생성 테스트"""

    def test_returns_excel_template(self, mapping_service):
        """엑셀 템플릿 바이트를 반환한다"""
        # When
        result = mapping_service.get_mapping_template()

        # Then
        assert isinstance(result, bytes)
        assert len(result) > 0

    def test_template_has_correct_headers(self, mapping_service):
        """템플릿에 올바른 헤더가 포함된다"""
        try:
            from openpyxl import load_workbook
        except ImportError:
            pytest.skip("openpyxl이 설치되지 않았습니다")

        # When
        template_bytes = mapping_service.get_mapping_template()
        wb = load_workbook(BytesIO(template_bytes))
        ws = wb.active

        # Then
        headers = [ws.cell(row=1, column=i).value for i in range(1, 8)]
        # 헤더에 필수 표시(*)가 포함되어 있으므로 부분 일치 확인
        headers_str = " ".join(headers)
        assert "자산코드" in headers_str
        assert "위협코드" in headers_str
        assert "취약점코드" in headers_str


# =============================================================================
# API 통합 테스트
# =============================================================================

class TestAssetRiskMappingAPI:
    """자산-위험 매핑 API 통합 테스트"""

    def test_get_recommended_threats_endpoint(
        self,
        client,
        auth_headers,
        asset_type_server,
        server_threats,
    ):
        """GET /api/v1/asset-risk-mapping/recommended-threats/{asset_type_id} 엔드포인트 테스트"""
        # When
        response = client.get(
            f"/api/v1/asset-risk-mapping/recommended-threats/{asset_type_server.id}",
            headers=auth_headers,
        )

        # Then
        assert response.status_code == 200
        data = response.json()
        assert "threats" in data
        assert len(data["threats"]) == 3

    def test_get_recommended_vulnerabilities_endpoint(
        self,
        client,
        auth_headers,
        asset_type_server,
        server_vulnerabilities,
    ):
        """GET /api/v1/asset-risk-mapping/recommended-vulnerabilities/{asset_type_id} 엔드포인트 테스트"""
        # When
        response = client.get(
            f"/api/v1/asset-risk-mapping/recommended-vulnerabilities/{asset_type_server.id}",
            headers=auth_headers,
        )

        # Then
        assert response.status_code == 200
        data = response.json()
        assert "vulnerabilities" in data

    def test_validate_mapping_endpoint(
        self,
        client,
        auth_headers,
        sample_asset,
        server_threats,
        server_vulnerabilities,
    ):
        """POST /api/v1/asset-risk-mapping/validate 엔드포인트 테스트"""
        # Given
        payload = {
            "asset_id": sample_asset.id,
            "threat_id": server_threats[0].id,
            "vulnerability_id": server_vulnerabilities[0].id,
        }

        # When
        response = client.post(
            "/api/v1/asset-risk-mapping/validate",
            json=payload,
            headers=auth_headers,
        )

        # Then
        assert response.status_code == 200
        data = response.json()
        assert "is_valid" in data

    def test_bulk_import_endpoint(
        self,
        client,
        auth_headers,
        sample_asset,
        sample_scenario,
        server_threats,
        server_vulnerabilities,
    ):
        """POST /api/v1/asset-risk-mapping/bulk-import 엔드포인트 테스트"""
        # Given
        excel_content = TestBulkCreateMappingsFromExcel()._create_valid_excel_file(
            sample_asset.asset_code,
            server_threats[0].code,
            server_vulnerabilities[0].code,
        )

        # When
        response = client.post(
            f"/api/v1/asset-risk-mapping/bulk-import?scenario_id={sample_scenario.id}",
            files={"file": ("mappings.xlsx", excel_content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            headers=auth_headers,
        )

        # Then
        assert response.status_code == 200
        data = response.json()
        assert "total" in data
        assert "success" in data

    def test_get_template_endpoint(self, client, auth_headers):
        """GET /api/v1/asset-risk-mapping/template 엔드포인트 테스트"""
        # When
        response = client.get(
            "/api/v1/asset-risk-mapping/template",
            headers=auth_headers,
        )

        # Then
        assert response.status_code == 200
        assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in response.headers.get("content-type", "")
