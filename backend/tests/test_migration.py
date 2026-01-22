"""
엑셀 데이터 마이그레이션 테스트
TDD - RED 단계: 테스트 먼저 작성
"""
import io
import pytest
from unittest.mock import patch, MagicMock
from datetime import datetime
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session
from openpyxl import Workbook

from app.services.migration_service import MigrationService, MigrationResult, MigrationError
from app.models.control import ControlDomain, ControlCategory, ControlItem
from app.models.evidence import Evidence


@pytest.fixture
def test_control_for_migration(db: Session) -> ControlItem:
    """마이그레이션 테스트용 통제항목 생성"""
    # 도메인 생성
    domain = db.query(ControlDomain).filter(ControlDomain.code == "M1").first()
    if not domain:
        domain = ControlDomain(
            code="M1",
            name="마이그레이션 테스트 도메인",
            sort_order=100,
        )
        db.add(domain)
        db.commit()
        db.refresh(domain)

    # 카테고리 생성
    category = db.query(ControlCategory).filter(ControlCategory.code == "M1.1").first()
    if not category:
        category = ControlCategory(
            domain_id=domain.id,
            code="M1.1",
            name="마이그레이션 테스트 카테고리",
            sort_order=1,
        )
        db.add(category)
        db.commit()
        db.refresh(category)

    # 통제항목 생성
    item = db.query(ControlItem).filter(ControlItem.code == "M1.1.1").first()
    if not item:
        item = ControlItem(
            category_id=category.id,
            code="M1.1.1",
            title="마이그레이션 테스트 항목",
            description="테스트용",
            is_required=True,
            sort_order=1,
        )
        db.add(item)
        db.commit()
        db.refresh(item)

    return item


def create_test_excel(data: list, headers: list = None) -> bytes:
    """테스트용 엑셀 파일 생성"""
    wb = Workbook()
    ws = wb.active

    # 기본 헤더
    if headers is None:
        headers = ["title", "description", "control_code", "valid_from", "valid_until", "author"]

    # 헤더 작성
    for col, header in enumerate(headers, start=1):
        ws.cell(row=1, column=col, value=header)

    # 데이터 작성
    for row_idx, row_data in enumerate(data, start=2):
        for col_idx, value in enumerate(row_data, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)

    # BytesIO로 저장
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.read()


class TestMigrationService:
    """MigrationService 단위 테스트"""

    def test_migration_error_to_dict(self):
        """MigrationError to_dict 변환"""
        error = MigrationError(row=5, column="title", message="필수값 누락", value=None)
        result = error.to_dict()

        assert result["row"] == 5
        assert result["column"] == "title"
        assert result["message"] == "필수값 누락"
        assert result["value"] is None

    def test_migration_result_add_error(self):
        """MigrationResult 오류 추가"""
        result = MigrationResult()
        result.add_error(row=3, column="control_code", message="존재하지 않는 코드")

        assert result.error_count == 1
        assert len(result.errors) == 1
        assert result.errors[0].row == 3

    def test_preview_success(
        self,
        db: Session,
        test_control_for_migration: ControlItem,
    ):
        """미리보기 성공"""
        # 테스트 데이터
        data = [
            ["테스트 증적 1", "설명 1", "M1.1.1", "2024-01-01", "2024-12-31", "홍길동"],
            ["테스트 증적 2", "설명 2", None, None, None, "김철수"],
        ]
        file_content = create_test_excel(data)

        service = MigrationService(db)
        result = service.preview(file_content)

        assert result.success_count == 2
        assert result.error_count == 0
        assert len(result.preview_data) == 2
        assert result.preview_data[0]["title"] == "테스트 증적 1"

    def test_preview_missing_title(self, db: Session):
        """제목 누락 시 오류"""
        data = [
            [None, "설명만 있음", None, None, None, None],
        ]
        file_content = create_test_excel(data)

        service = MigrationService(db)
        result = service.preview(file_content)

        assert result.error_count >= 1
        assert any(e.column == "title" for e in result.errors)

    def test_preview_invalid_control_code(
        self,
        db: Session,
        test_control_for_migration: ControlItem,
    ):
        """존재하지 않는 통제항목 코드"""
        data = [
            ["테스트 증적", "설명", "INVALID_CODE", None, None, None],
        ]
        file_content = create_test_excel(data)

        service = MigrationService(db)
        result = service.preview(file_content)

        assert result.error_count >= 1
        assert any(e.column == "control_code" for e in result.errors)

    def test_preview_invalid_date_format(self, db: Session):
        """잘못된 날짜 형식"""
        data = [
            ["테스트 증적", "설명", None, "invalid-date", None, None],
        ]
        file_content = create_test_excel(data)

        service = MigrationService(db)
        result = service.preview(file_content)

        assert result.error_count >= 1
        assert any(e.column == "valid_from" for e in result.errors)

    def test_preview_missing_required_column(self, db: Session):
        """필수 컬럼 누락"""
        # title 컬럼 없이 생성
        wb = Workbook()
        ws = wb.active
        ws.cell(row=1, column=1, value="description")
        ws.cell(row=2, column=1, value="설명만")

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        file_content = output.read()

        service = MigrationService(db)
        result = service.preview(file_content)

        assert result.error_count >= 1
        assert any("필수 컬럼 누락" in e.message for e in result.errors)

    def test_execute_success(
        self,
        db: Session,
        test_control_for_migration: ControlItem,
        test_admin_user,
    ):
        """마이그레이션 실행 성공"""
        data = [
            ["실행 테스트 증적", "설명", "M1.1.1", "2024-01-01", "2024-12-31", "홍길동"],
        ]
        file_content = create_test_excel(data)

        service = MigrationService(db)
        result = service.execute(
            file_content=file_content,
            uploader_id=test_admin_user.id,
        )

        assert result.success_count == 1
        assert len(result.created_evidences) == 1

        # DB에서 확인
        evidence = db.query(Evidence).filter(Evidence.id == result.created_evidences[0]).first()
        assert evidence is not None
        assert evidence.title == "실행 테스트 증적"

    def test_execute_skip_errors(
        self,
        db: Session,
        test_admin_user,
    ):
        """오류 행 건너뛰기"""
        data = [
            ["정상 증적", "설명", None, None, None, None],
            [None, "제목 없음", None, None, None, None],  # 오류
            ["또 다른 정상 증적", "설명", None, None, None, None],
        ]
        file_content = create_test_excel(data)

        service = MigrationService(db)
        result = service.execute(
            file_content=file_content,
            uploader_id=test_admin_user.id,
            skip_errors=True,
        )

        # 오류 건너뛰고 2개 생성
        assert result.success_count == 2
        assert result.error_count >= 1


class TestMigrationPreviewAPI:
    """POST /migration/preview 테스트"""

    def test_preview_success(
        self,
        client: TestClient,
        admin_auth_headers: dict,
        test_control_for_migration: ControlItem,
    ):
        """미리보기 API 성공"""
        data = [
            ["API 테스트 증적", "설명", "M1.1.1", None, None, None],
        ]
        file_content = create_test_excel(data)

        files = {"file": ("test.xlsx", io.BytesIO(file_content), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}

        response = client.post(
            "/api/v1/migration/preview",
            headers=admin_auth_headers,
            files=files,
        )

        assert response.status_code == 200
        result = response.json()
        assert "migration_id" in result
        assert result["success_count"] >= 1

    def test_preview_invalid_file_type(
        self,
        client: TestClient,
        admin_auth_headers: dict,
    ):
        """잘못된 파일 형식"""
        files = {"file": ("test.txt", io.BytesIO(b"not excel"), "text/plain")}

        response = client.post(
            "/api/v1/migration/preview",
            headers=admin_auth_headers,
            files=files,
        )

        assert response.status_code == 400

    def test_preview_unauthorized(
        self,
        client: TestClient,
        auth_headers: dict,  # 일반 사용자
    ):
        """권한 없는 사용자"""
        data = [["테스트", None, None, None, None, None]]
        file_content = create_test_excel(data)
        files = {"file": ("test.xlsx", io.BytesIO(file_content), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}

        response = client.post(
            "/api/v1/migration/preview",
            headers=auth_headers,
            files=files,
        )

        assert response.status_code == 403


class TestMigrationExecuteAPI:
    """POST /migration/execute 테스트"""

    def test_execute_success(
        self,
        client: TestClient,
        admin_auth_headers: dict,
        test_control_for_migration: ControlItem,
    ):
        """마이그레이션 실행 성공"""
        # 먼저 미리보기
        data = [
            ["실행 API 테스트", "설명", "M1.1.1", None, None, None],
        ]
        file_content = create_test_excel(data)
        files = {"file": ("test.xlsx", io.BytesIO(file_content), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}

        preview_response = client.post(
            "/api/v1/migration/preview",
            headers=admin_auth_headers,
            files=files,
        )
        migration_id = preview_response.json()["migration_id"]

        # 실행
        response = client.post(
            "/api/v1/migration/execute",
            headers=admin_auth_headers,
            data={
                "migration_id": migration_id,
                "skip_errors": "false",
            },
        )

        assert response.status_code == 200
        result = response.json()
        assert result["success_count"] >= 1
        assert len(result["created_evidence_ids"]) >= 1

    def test_execute_not_found(
        self,
        client: TestClient,
        admin_auth_headers: dict,
    ):
        """존재하지 않는 migration_id"""
        response = client.post(
            "/api/v1/migration/execute",
            headers=admin_auth_headers,
            data={
                "migration_id": "non-existent-id",
                "skip_errors": "false",
            },
        )

        assert response.status_code == 404


class TestMigrationErrorsAPI:
    """GET /migration/{id}/errors 테스트"""

    def test_get_errors_success(
        self,
        client: TestClient,
        admin_auth_headers: dict,
    ):
        """오류 리포트 조회"""
        # 오류가 있는 데이터로 미리보기
        data = [
            [None, "제목 없음", None, None, None, None],  # 오류
            ["정상", None, None, None, None, None],
        ]
        file_content = create_test_excel(data)
        files = {"file": ("test.xlsx", io.BytesIO(file_content), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}

        preview_response = client.post(
            "/api/v1/migration/preview",
            headers=admin_auth_headers,
            files=files,
        )
        migration_id = preview_response.json()["migration_id"]

        # 오류 조회
        response = client.get(
            f"/api/v1/migration/{migration_id}/errors",
            headers=admin_auth_headers,
        )

        assert response.status_code == 200
        result = response.json()
        assert "errors" in result
        assert "total" in result

    def test_get_errors_not_found(
        self,
        client: TestClient,
        admin_auth_headers: dict,
    ):
        """존재하지 않는 마이그레이션 오류 조회"""
        response = client.get(
            "/api/v1/migration/non-existent-id/errors",
            headers=admin_auth_headers,
        )

        assert response.status_code == 404
