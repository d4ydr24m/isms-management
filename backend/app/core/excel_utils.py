"""Excel 생성 공용 유틸리티"""
from openpyxl.styles import PatternFill
from openpyxl.worksheet.worksheet import Worksheet


def apply_banded_rows(
    ws: Worksheet,
    start_row: int,
    end_row: int,
    num_cols: int,
    band_color: str = "F7F9FC",
) -> None:
    """홀수 인덱스 행에 옅은 배경색 적용 (줄무늬 효과).

    - 기존 배경색이 있는 셀은 건너뜀 (예시 행 등).
    - 짝수 인덱스 행(첫 번째 데이터 행 포함)은 흰색 그대로.
    """
    band_fill = PatternFill(start_color=band_color, end_color=band_color, fill_type="solid")
    for row_num in range(start_row, end_row + 1):
        if (row_num - start_row) % 2 == 1:
            for col_num in range(1, num_cols + 1):
                cell = ws.cell(row=row_num, column=col_num)
                if cell.fill.start_color.rgb in (None, "00000000"):
                    cell.fill = band_fill
