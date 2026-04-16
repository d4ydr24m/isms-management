"""
자산-위협-취약점 3-way 매핑 서비스
Phase 2: Section 5.1

기능:
- 5.1.1: 자산 유형별 위협/취약점 자동 추천
- 5.1.2: 매핑 유효성 검증 로직
- 5.1.3: 대량 매핑 기능 (엑셀 임포트)
"""
from datetime import datetime, timezone
from io import BytesIO
from typing import Any, Dict, List, Optional

from sqlalchemy import and_
from sqlalchemy.orm import Session

from app.models.asset import Asset, AssetType
from app.models.risk import (
    Threat,
    Vulnerability,
    AssetTypeThreat,
    RiskScenario,
    RiskAssessment,
)


def utc_now():
    """UTC 현재 시간 반환"""
    return datetime.now(timezone.utc)


class AssetRiskMappingService:
    """
    자산-위험 매핑 서비스

    자산 유형에 따른 위협/취약점 추천, 매핑 유효성 검증,
    엑셀 임포트를 통한 대량 매핑 기능을 제공합니다.
    """

    def __init__(self, db: Session):
        self.db = db

    # =========================================================================
    # 5.1.1: 자산 유형별 위협/취약점 자동 추천
    # =========================================================================

    def get_recommended_threats(self, asset_type_id: int) -> List[Threat]:
        """
        자산 유형에 대해 관련 위협을 추천합니다.

        AssetTypeThreat 매핑 테이블을 통해 해당 자산 유형과
        연결된 위협 목록을 관련성 점수 순으로 반환합니다.

        Args:
            asset_type_id: 자산 유형 ID

        Returns:
            관련성 점수 내림차순으로 정렬된 Threat 객체 목록
        """
        # 자산 유형 존재 확인
        asset_type = self.db.query(AssetType).filter(
            AssetType.id == asset_type_id
        ).first()

        if not asset_type:
            return []

        # AssetTypeThreat 매핑을 통해 관련 위협 조회
        mappings = (
            self.db.query(AssetTypeThreat)
            .filter(
                AssetTypeThreat.asset_type_id == asset_type_id,
                AssetTypeThreat.is_active == True,
            )
            .order_by(AssetTypeThreat.relevance_score.desc())
            .all()
        )

        # 위협 객체 수집 (활성 상태만)
        threats = []
        for mapping in mappings:
            threat = self.db.query(Threat).filter(
                Threat.id == mapping.threat_id,
                Threat.is_active == True,
            ).first()
            if threat:
                threats.append(threat)

        return threats

    def get_recommended_vulnerabilities(
        self, asset_type_id: int
    ) -> List[Vulnerability]:
        """
        자산 유형에 대해 관련 취약점을 추천합니다.

        현재는 모든 활성 취약점을 반환합니다.
        향후 자산 유형-취약점 매핑 테이블 추가 시 확장 가능합니다.

        Args:
            asset_type_id: 자산 유형 ID

        Returns:
            활성 상태의 Vulnerability 객체 목록
        """
        # 자산 유형 존재 확인
        asset_type = self.db.query(AssetType).filter(
            AssetType.id == asset_type_id
        ).first()

        if not asset_type:
            return []

        # 현재는 모든 활성 취약점 반환
        # 향후 AssetTypeVulnerability 매핑 추가 시 해당 로직으로 대체
        vulnerabilities = (
            self.db.query(Vulnerability)
            .filter(Vulnerability.is_active == True)
            .order_by(Vulnerability.vulnerability_level.desc(), Vulnerability.code)
            .all()
        )

        return vulnerabilities

    # =========================================================================
    # 5.1.2: 매핑 유효성 검증 로직
    # =========================================================================

    def validate_mapping(
        self,
        asset_id: int,
        threat_id: int,
        vulnerability_id: int,
    ) -> Dict[str, Any]:
        """
        자산-위협-취약점 매핑의 유효성을 검증합니다.

        검증 항목:
        - 각 엔티티 존재 여부
        - 활성 상태 확인
        - 자산 유형과 위협의 관련성 (경고 수준)

        Args:
            asset_id: 자산 ID
            threat_id: 위협 ID
            vulnerability_id: 취약점 ID

        Returns:
            Dict containing:
            - is_valid: 유효성 여부 (bool)
            - errors: 에러 메시지 목록 (list)
            - warnings: 경고 메시지 목록 (list)
        """
        errors = []
        warnings = []

        # 1. 자산 검증
        asset = self.db.query(Asset).filter(Asset.id == asset_id).first()
        if not asset:
            errors.append(f"자산을 찾을 수 없습니다. (ID: {asset_id})")
        elif not asset.is_active:
            errors.append(f"비활성 상태의 자산입니다. (ID: {asset_id})")

        # 2. 위협 검증
        threat = self.db.query(Threat).filter(Threat.id == threat_id).first()
        if not threat:
            errors.append(f"위협을 찾을 수 없습니다. (ID: {threat_id})")
        elif not threat.is_active:
            errors.append(f"비활성 상태의 위협입니다. (ID: {threat_id})")

        # 3. 취약점 검증
        vulnerability = self.db.query(Vulnerability).filter(
            Vulnerability.id == vulnerability_id
        ).first()
        if not vulnerability:
            errors.append(f"취약점을 찾을 수 없습니다. (ID: {vulnerability_id})")
        elif not vulnerability.is_active:
            errors.append(f"비활성 상태의 취약점입니다. (ID: {vulnerability_id})")

        # 4. 자산 유형-위협 관련성 검증 (에러가 없는 경우만)
        if not errors and asset and threat:
            is_related = (
                self.db.query(AssetTypeThreat)
                .filter(
                    AssetTypeThreat.asset_type_id == asset.asset_type_id,
                    AssetTypeThreat.threat_id == threat_id,
                    AssetTypeThreat.is_active == True,
                )
                .first()
            )

            if not is_related:
                warnings.append(
                    f"선택한 위협({threat.name})은 해당 자산 유형과 일반적으로 연관되지 않습니다. "
                    "계속 진행할 수 있지만 검토를 권장합니다."
                )

        return {
            "is_valid": len(errors) == 0,
            "errors": errors,
            "warnings": warnings,
        }

    def validate_mapping_batch(
        self, mappings: List[Dict[str, int]]
    ) -> Dict[str, Any]:
        """
        여러 매핑을 일괄 검증합니다.

        Args:
            mappings: 매핑 정보 목록
                [{"asset_id": int, "threat_id": int, "vulnerability_id": int}, ...]

        Returns:
            Dict containing:
            - total: 전체 개수
            - valid_count: 유효한 매핑 개수
            - invalid_count: 유효하지 않은 매핑 개수
            - results: 개별 검증 결과 목록
        """
        results = []
        valid_count = 0
        invalid_count = 0

        for idx, mapping in enumerate(mappings):
            result = self.validate_mapping(
                asset_id=mapping.get("asset_id"),
                threat_id=mapping.get("threat_id"),
                vulnerability_id=mapping.get("vulnerability_id"),
            )
            result["row"] = idx + 1
            results.append(result)

            if result["is_valid"]:
                valid_count += 1
            else:
                invalid_count += 1

        return {
            "total": len(mappings),
            "valid_count": valid_count,
            "invalid_count": invalid_count,
            "results": results,
        }

    # =========================================================================
    # 5.1.3: 대량 매핑 기능 (엑셀 임포트)
    # =========================================================================

    def bulk_create_mappings_from_excel(
        self,
        file_content: bytes,
        scenario_id: int,
        user_id: int,
    ) -> Dict[str, Any]:
        """
        엑셀 파일에서 매핑 정보를 읽어 대량으로 위험 평가를 생성합니다.

        예상 컬럼:
        - 자산코드: Asset.asset_code
        - 위협코드: Threat.code
        - 취약점코드: Vulnerability.code
        - 자산가치: 1-3 (선택)
        - 위협등급: 1-3 (선택)
        - 취약점등급: 1-3 (선택)
        - 비고: 평가 의견 (선택)

        Args:
            file_content: 엑셀 파일 바이트
            scenario_id: 위험 평가 시나리오 ID
            user_id: 작업자 ID

        Returns:
            Dict containing:
            - total: 전체 행 수
            - success: 성공 개수
            - failed: 실패 개수
            - skipped: 건너뛴 개수 (중복)
            - errors: 에러 목록
        """
        try:
            from openpyxl import load_workbook
        except ImportError:
            return {
                "total": 0,
                "success": 0,
                "failed": 1,
                "skipped": 0,
                "errors": [{"row": 0, "error": "openpyxl 패키지가 필요합니다."}],
            }

        # 시나리오 존재 확인
        scenario = self.db.query(RiskScenario).filter(
            RiskScenario.id == scenario_id
        ).first()
        if not scenario:
            return {
                "total": 0,
                "success": 0,
                "failed": 1,
                "skipped": 0,
                "errors": [{"row": 0, "error": f"시나리오를 찾을 수 없습니다. (ID: {scenario_id})"}],
            }

        try:
            wb = load_workbook(BytesIO(file_content))
            ws = wb.active
        except Exception as e:
            return {
                "total": 0,
                "success": 0,
                "failed": 1,
                "skipped": 0,
                "errors": [{"row": 0, "error": f"엑셀 파일 읽기 실패: {str(e)}"}],
            }

        results = {
            "total": 0,
            "success": 0,
            "failed": 0,
            "skipped": 0,
            "errors": [],
        }

        # 헤더 행 건너뛰기
        for row_num in range(2, ws.max_row + 1):
            asset_code = ws.cell(row=row_num, column=1).value
            threat_code = ws.cell(row=row_num, column=2).value
            vuln_code = ws.cell(row=row_num, column=3).value

            # 빈 행 건너뛰기
            if not asset_code and not threat_code and not vuln_code:
                continue

            results["total"] += 1

            # 필수 값 검증
            if not asset_code or not threat_code or not vuln_code:
                results["failed"] += 1
                results["errors"].append({
                    "row": row_num,
                    "error": "자산코드, 위협코드, 취약점코드는 필수입니다.",
                })
                continue

            # 자산 조회
            asset = self.db.query(Asset).filter(
                Asset.asset_code == str(asset_code)
            ).first()
            if not asset:
                results["failed"] += 1
                results["errors"].append({
                    "row": row_num,
                    "error": f"자산을 찾을 수 없습니다: {asset_code}",
                })
                continue

            # 위협 조회
            threat = self.db.query(Threat).filter(
                Threat.code == str(threat_code)
            ).first()
            if not threat:
                results["failed"] += 1
                results["errors"].append({
                    "row": row_num,
                    "error": f"위협을 찾을 수 없습니다: {threat_code}",
                })
                continue

            # 취약점 조회
            vulnerability = self.db.query(Vulnerability).filter(
                Vulnerability.code == str(vuln_code)
            ).first()
            if not vulnerability:
                results["failed"] += 1
                results["errors"].append({
                    "row": row_num,
                    "error": f"취약점을 찾을 수 없습니다: {vuln_code}",
                })
                continue

            # 중복 확인
            existing = self.db.query(RiskAssessment).filter(
                and_(
                    RiskAssessment.scenario_id == scenario_id,
                    RiskAssessment.asset_id == asset.id,
                    RiskAssessment.threat_id == threat.id,
                    RiskAssessment.vulnerability_id == vulnerability.id,
                )
            ).first()

            if existing:
                results["skipped"] += 1
                continue

            # 선택적 값 읽기
            asset_value = ws.cell(row=row_num, column=4).value or 2
            threat_level = ws.cell(row=row_num, column=5).value or threat.threat_level
            vuln_level = ws.cell(row=row_num, column=6).value or vulnerability.vulnerability_level
            remarks = ws.cell(row=row_num, column=7).value

            # 값 범위 검증 (1-3)
            try:
                asset_value = max(1, min(3, int(asset_value)))
                threat_level = max(1, min(3, int(threat_level)))
                vuln_level = max(1, min(3, int(vuln_level)))
            except (ValueError, TypeError):
                asset_value = 2
                threat_level = threat.threat_level
                vuln_level = vulnerability.vulnerability_level

            # 위험 평가 생성
            try:
                assessment = RiskAssessment(
                    scenario_id=scenario_id,
                    asset_id=asset.id,
                    threat_id=threat.id,
                    vulnerability_id=vulnerability.id,
                    asset_value=asset_value,
                    threat_level=threat_level,
                    vulnerability_level=vuln_level,
                    evaluated_by=user_id,
                    evaluated_at=utc_now(),
                    remarks=str(remarks) if remarks else None,
                )
                self.db.add(assessment)
                self.db.flush()  # ID 생성을 위해
                results["success"] += 1
            except Exception as e:
                results["failed"] += 1
                results["errors"].append({
                    "row": row_num,
                    "error": f"위험 평가 생성 실패: {str(e)}",
                })

        # 커밋
        if results["success"] > 0:
            self.db.commit()

        return results

    def get_mapping_template(self) -> bytes:
        """
        대량 매핑용 엑셀 템플릿을 생성합니다.

        Returns:
            엑셀 파일 바이트
        """
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment
        except ImportError:
            raise ImportError("openpyxl 패키지가 필요합니다.")

        wb = Workbook()
        ws = wb.active
        ws.title = "매핑"

        # 헤더 스타일
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(
            start_color="4472C4", end_color="4472C4", fill_type="solid"
        )
        required_fill = PatternFill(
            start_color="FFC000", end_color="FFC000", fill_type="solid"
        )

        # 헤더 정의
        headers = [
            ("자산코드*", True),
            ("위협코드*", True),
            ("취약점코드*", True),
            ("자산가치 (1-3)", False),
            ("위협등급 (1-3)", False),
            ("취약점등급 (1-3)", False),
            ("비고", False),
        ]

        for col, (header, required) in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = required_fill if required else header_fill
            cell.alignment = Alignment(horizontal="center")

        # 예시 데이터
        ws.cell(row=2, column=1, value="AST-SRV-202501-001")
        ws.cell(row=2, column=2, value="T-SRV-001")
        ws.cell(row=2, column=3, value="V-SRV-001")
        ws.cell(row=2, column=4, value=2)
        ws.cell(row=2, column=5, value=2)
        ws.cell(row=2, column=6, value=2)
        ws.cell(row=2, column=7, value="예시 매핑")

        # 열 너비 조정
        column_widths = [25, 20, 20, 15, 15, 15, 30]
        for col, width in enumerate(column_widths, 1):
            ws.column_dimensions[chr(64 + col)].width = width

        # 헤더 자동 필터 및 틀 고정
        ws.auto_filter.ref = "A1:G1"
        ws.freeze_panes = "A2"

        # 바이트로 변환
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        return output.read()

    # =========================================================================
    # 추가 유틸리티 메서드
    # =========================================================================

    def get_asset_type_threat_mappings(
        self, asset_type_id: int
    ) -> List[Dict[str, Any]]:
        """
        자산 유형의 위협 매핑 상세 정보를 반환합니다.

        Args:
            asset_type_id: 자산 유형 ID

        Returns:
            매핑 정보 목록 (위협 정보 + 관련성 점수)
        """
        mappings = (
            self.db.query(AssetTypeThreat)
            .filter(
                AssetTypeThreat.asset_type_id == asset_type_id,
                AssetTypeThreat.is_active == True,
            )
            .order_by(AssetTypeThreat.relevance_score.desc())
            .all()
        )

        result = []
        for mapping in mappings:
            threat = self.db.query(Threat).filter(
                Threat.id == mapping.threat_id
            ).first()
            if threat:
                result.append({
                    "threat_id": threat.id,
                    "threat_code": threat.code,
                    "threat_name": threat.name,
                    "threat_level": threat.threat_level,
                    "relevance_score": mapping.relevance_score,
                })

        return result

    def create_asset_type_threat_mapping(
        self,
        asset_type_id: int,
        threat_id: int,
        relevance_score: float = 0.8,
    ) -> AssetTypeThreat:
        """
        자산 유형-위협 매핑을 생성합니다.

        Args:
            asset_type_id: 자산 유형 ID
            threat_id: 위협 ID
            relevance_score: 관련성 점수 (0.0 ~ 1.0)

        Returns:
            생성된 AssetTypeThreat 객체
        """
        # 중복 확인
        existing = self.db.query(AssetTypeThreat).filter(
            AssetTypeThreat.asset_type_id == asset_type_id,
            AssetTypeThreat.threat_id == threat_id,
        ).first()

        if existing:
            # 기존 매핑 활성화
            existing.is_active = True
            existing.relevance_score = relevance_score
            self.db.commit()
            self.db.refresh(existing)
            return existing

        # 새 매핑 생성
        mapping = AssetTypeThreat(
            asset_type_id=asset_type_id,
            threat_id=threat_id,
            relevance_score=max(0.0, min(1.0, relevance_score)),
            is_active=True,
        )
        self.db.add(mapping)
        self.db.commit()
        self.db.refresh(mapping)
        return mapping
