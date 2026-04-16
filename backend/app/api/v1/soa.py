"""
SOA(적용성 보고서) API 라우터
/api/v1/soa
Phase 2: FR-606
"""
from typing import List

from fastapi import APIRouter, Depends, HTTPException, status
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from io import BytesIO

from app.core.deps import get_db, require_permission
from app.models.control import ControlItem
from app.models.user import User
from app.api.v1.system_settings import get_certification_type
from app.schemas.risk import (
    SOARecordResponse,
    SOARecordUpdate,
    SOARecordList,
    SOAExportRequest,
)
from app.services.risk_service import RiskService


router = APIRouter()


# =============================================================================
# 헬퍼 함수
# =============================================================================

def get_risk_service(db: Session = Depends(get_db)) -> RiskService:
    """RiskService 의존성"""
    return RiskService(db)


IMPLEMENTATION_STATUS_NAMES = {
    "fully_implemented": "완전 구현",
    "partially_implemented": "부분 구현",
    "planned": "계획",
    "not_implemented": "미구현",
    "not_applicable": "해당없음",
}


def soa_record_to_response(record) -> SOARecordResponse:
    """SOARecord 모델을 응답으로 변환"""
    control = record.control_item
    return SOARecordResponse(
        id=record.id,
        control_item_id=record.control_item_id,
        control_code=control.code if control else None,
        control_title=control.title if control else None,
        control_description=control.description if control else None,
        is_applicable=record.is_applicable,
        exclusion_reason=record.exclusion_reason,
        implementation_status=record.implementation_status,
        implementation_status_name=IMPLEMENTATION_STATUS_NAMES.get(record.implementation_status),
        implementation_evidence=record.implementation_evidence,
        related_assets=record.related_assets,
        related_risks=record.related_risks,
        remarks=record.remarks,
        created_at=record.created_at,
        updated_at=record.updated_at,
    )


# =============================================================================
# 4.8: SOA API (FR-606)
# =============================================================================

@router.get("", response_model=SOARecordList)
def get_soa_list(
    db: Session = Depends(get_db),
    service: RiskService = Depends(get_risk_service),
    current_user: User = Depends(require_permission("risk:read")),
) -> SOARecordList:
    """
    SOA 목록 조회

    인증 유형(ISMS/ISMS-P)에 따라 통제항목의 적용 여부 및 구현 상태를 조회합니다.
    """
    cert_type = get_certification_type(db)
    items, total = service.get_soa_records()

    # ISMS 모드에서는 개인정보 관련 통제항목 제외
    if cert_type == "ISMS":
        items = [item for item in items if not (item.control_item and item.control_item.is_personal_info)]
        total = len(items)

    applicable_count = sum(1 for item in items if item.is_applicable)
    not_applicable_count = total - applicable_count

    return SOARecordList(
        items=[soa_record_to_response(item) for item in items],
        total=total,
        applicable_count=applicable_count,
        not_applicable_count=not_applicable_count,
    )


@router.put("/{control_id}", response_model=SOARecordResponse)
def update_soa_record(
    control_id: int,
    data: SOARecordUpdate,
    service: RiskService = Depends(get_risk_service),
    current_user: User = Depends(require_permission("risk:update")),
) -> SOARecordResponse:
    """
    SOA 레코드 수정

    - **is_applicable**: 적용 여부
    - **exclusion_reason**: 제외 사유 (미적용 시)
    - **implementation_status**: 구현 상태
    - **implementation_evidence**: 구현 증적 설명
    - **related_assets**: 관련 자산 목록
    - **related_risks**: 관련 위험 목록
    """
    record = service.update_soa_record(
        control_item_id=control_id,
        **data.model_dump(exclude_unset=True),
    )
    return soa_record_to_response(record)


@router.post("/generate")
def generate_soa(
    service: RiskService = Depends(get_risk_service),
    current_user: User = Depends(require_permission("risk:create")),
):
    """
    SOA 자동 생성

    모든 통제항목에 대해 SOA 레코드를 자동 생성합니다.
    이미 존재하는 레코드는 건너뜁니다.
    """
    count = service.generate_soa()
    return {"created_count": count, "message": f"{count}개의 SOA 레코드가 생성되었습니다."}


@router.get("/export")
def export_soa(
    format: str = "excel",
    template_type: str = "isms_p",
    service: RiskService = Depends(get_risk_service),
    current_user: User = Depends(require_permission("risk:read")),
):
    """
    SOA 내보내기

    - **format**: 내보내기 형식 (excel/word)
    - **template_type**: 템플릿 유형 (isms_p/iso27001)
    """
    items, _ = service.get_soa_records()

    if format == "excel":
        content = _generate_soa_excel(items, template_type)
        media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        filename = f"SOA_{template_type}.xlsx"
    elif format == "word":
        content = _generate_soa_word(items, template_type)
        media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        filename = f"SOA_{template_type}.docx"
    else:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="지원하지 않는 형식입니다.",
        )

    return StreamingResponse(
        BytesIO(content),
        media_type=media_type,
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


def _generate_soa_excel(items, template_type: str) -> bytes:
    """SOA 엑셀 생성"""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    except ImportError:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="openpyxl 패키지가 필요합니다.",
        )

    wb = Workbook()
    ws = wb.active
    ws.title = "적용성보고서(SOA)"

    # 스타일 정의
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    # 헤더
    headers = [
        "통제항목 코드",
        "통제항목명",
        "적용 여부",
        "제외 사유",
        "구현 상태",
        "구현 증적",
        "관련 자산",
        "관련 위험",
        "비고",
    ]

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    # 데이터
    for row, record in enumerate(items, 2):
        control = record.control_item
        ws.cell(row=row, column=1, value=control.code if control else "").border = thin_border
        ws.cell(row=row, column=2, value=control.title if control else "").border = thin_border
        ws.cell(row=row, column=3, value="적용" if record.is_applicable else "미적용").border = thin_border
        ws.cell(row=row, column=4, value=record.exclusion_reason or "").border = thin_border
        ws.cell(row=row, column=5, value=IMPLEMENTATION_STATUS_NAMES.get(record.implementation_status, "")).border = thin_border
        ws.cell(row=row, column=6, value=record.implementation_evidence or "").border = thin_border
        ws.cell(row=row, column=7, value=record.related_assets or "").border = thin_border
        ws.cell(row=row, column=8, value=record.related_risks or "").border = thin_border
        ws.cell(row=row, column=9, value=record.remarks or "").border = thin_border

    # 열 너비 조정
    column_widths = [15, 40, 10, 30, 15, 40, 30, 30, 30]
    for col, width in enumerate(column_widths, 1):
        ws.column_dimensions[chr(64 + col)].width = width

    ws.auto_filter.ref = "A1:I1"

    # 바이트로 반환
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output.read()


def _generate_soa_word(items, template_type: str) -> bytes:
    """SOA Word 문서 생성"""
    try:
        from docx import Document
        from docx.shared import Inches, Pt
        from docx.enum.table import WD_TABLE_ALIGNMENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="python-docx 패키지가 필요합니다.",
        )

    doc = Document()

    # 제목
    title = doc.add_heading("적용성 보고서 (Statement of Applicability)", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 템플릿 유형에 따른 부제목
    if template_type == "isms_p":
        doc.add_paragraph("ISMS-P 인증기준 적용성 보고서")
    else:
        doc.add_paragraph("ISO27001 적용성 보고서")

    doc.add_paragraph()

    # 통계 요약
    applicable = sum(1 for item in items if item.is_applicable)
    not_applicable = len(items) - applicable
    doc.add_paragraph(f"총 통제항목: {len(items)}개")
    doc.add_paragraph(f"적용: {applicable}개 / 미적용: {not_applicable}개")
    doc.add_paragraph()

    # 테이블 생성
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # 헤더
    header_cells = table.rows[0].cells
    headers = ["통제항목", "적용", "구현 상태", "증적", "비고"]
    for i, header in enumerate(headers):
        header_cells[i].text = header
        header_cells[i].paragraphs[0].runs[0].bold = True

    # 데이터
    for record in items:
        control = record.control_item
        row_cells = table.add_row().cells
        row_cells[0].text = f"{control.code}\n{control.title}" if control else ""
        row_cells[1].text = "O" if record.is_applicable else "X"
        row_cells[2].text = IMPLEMENTATION_STATUS_NAMES.get(record.implementation_status, "")
        row_cells[3].text = record.implementation_evidence or ""
        row_cells[4].text = record.remarks or ""

    # 바이트로 반환
    output = BytesIO()
    doc.save(output)
    output.seek(0)
    return output.read()
